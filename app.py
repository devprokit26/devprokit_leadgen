"""app.py - simple web UI: login -> search. Chalao: python app.py  phir browser mein http://127.0.0.1:5055"""
import contextlib, io, os, random, socket, threading, time, webbrowser
from functools import wraps

from flask import Flask, jsonify, redirect, request, send_file, session
from openpyxl import Workbook

import json, httpx
import config, email_finder, main, places, web_search
from excel_store import ExcelStore, COLUMNS

app = Flask(__name__)
app.secret_key = getattr(config, "APP_SECRET", "secret")

# Sirf gloves: har category ke keywords (English + Germany ke liye German)
CATS = {
    "sports": {"en": ["sports gloves supplier", "sports gloves wholesaler"], "de": ["Sporthandschuhe Grosshandel"],
               "words": ["sports glove", "goalkeeper", "cycling glove", "training glove", "sporthandschuh"]},
    "boxing": {"en": ["boxing gloves supplier", "boxing gloves wholesaler"], "de": ["Boxhandschuhe Grosshandel"],
               "words": ["boxing", "boxhandschuh", "boxen", "mma glove"]},
    "medical": {"en": ["medical gloves distributor", "nitrile examination gloves supplier"], "de": ["Untersuchungshandschuhe Grosshandel"],
                "words": ["medical", "nitrile", "examination", "untersuchung", "medizin"]},
    "surgical": {"en": ["surgical gloves supplier", "sterile surgical gloves distributor"], "de": ["OP-Handschuhe Lieferant"],
                 "words": ["surgical", "sterile", "operation", "op-handschuh", "chirurg"]},
}
BASE_WORDS = ["glove", "handschuh", "guantes", "guanti", "gants", "handschoen"]
BASE_EXCLUDE = ["baumarkt", "spielzeug", "toys", "nutrition", "bootsladen", "softair"]
NON_SPORT_EXCLUDE = ["golf", "sport", "fitness", "fight", "boxen", "boxing", "motorrad", "dainese", "segeln", "angeln", "reiten"]

JOB = {"saved": 0, "phase": "", "running": False, "stop": False, "log": [], "done": 0, "total": 0, "leads": [], "msg": ""}


class Log(io.TextIOBase):
    def write(self, s):
        for line in s.splitlines():
            if line.strip():
                JOB["log"].append(line.strip()[:200])
        return len(s)


def make_queries(country, places_, cats):
    qs = []
    for place, city in places_:
        for c in cats:
            kws = CATS[c]["en"] + (CATS[c]["de"] if country == "Germany" else [])
            qs += [(f"{k} {place}, {country}", city) for k in kws]
    return qs


def open_store():
    """leads.xlsx kharab (zip nahi) ho to usay backup naam se side par rakh kar nayi file banata hai."""
    path = config.EXCEL_FILE
    try:
        return ExcelStore(path)
    except Exception as e:
        print("leads.xlsx kharab hai (%s)" % e)
        if os.path.exists(path):
            os.replace(path, path.replace(".xlsx", "_broken_%d.xlsx" % int(time.time())))
        return ExcelStore(path)


def run_job(country, places_, cats, max_per):
    JOB.update(running=True, stop=False, log=[], done=0, leads=[], msg="", saved=0, phase="search")
    try:
        with contextlib.redirect_stdout(Log()):
            store = open_store()
            start = store.next_id
            words = BASE_WORDS + [w for c in cats for w in CATS[c]["words"]]
            exclude = BASE_EXCLUDE + ([] if ("sports" in cats or "boxing" in cats) else NON_SPORT_EXCLUDE)
            accept = main.make_relevance_checker(words, exclude)
            qs = make_queries(country, places_, cats)
            JOB["total"] = len(qs)
            for i, (q, city) in enumerate(qs, 1):
                if JOB["stop"] or store.remaining_today() <= 0:
                    break
                st = web_search.run(q, max_per, store, country=country, city=city, accept=accept)
                JOB["done"] = i
                JOB["saved"] += st["saved"]
                if st["blocked"]:
                    JOB["msg"] = "Search engine ne rok diya, 30-60 min baad dobara try karo."
                    break
                time.sleep(random.uniform(config.SEARCH_MIN_DELAY, config.SEARCH_MAX_DELAY))
            JOB["phase"] = "contacts"
            print("Email / owner / social nikal rahe hain...")
            email_finder.process_file(store, workers=10)
            keep = ["business_name", "phone", "email", "website", "city", "country", "owner_name", "owner_title",
                    "facebook", "instagram", "linkedin", "twitter", "youtube", "tiktok"]
            for r in range(2, store.ws.max_row + 1):
                d = dict(zip(COLUMNS, [c.value for c in store.ws[r]]))
                if d.get("id") and int(d["id"]) >= start:
                    JOB["leads"].append({k: d.get(k) or "" for k in keep})
    except Exception as e:
        JOB["msg"] = f"Error: {e}"
    JOB["running"] = False


def login_required(f):
    @wraps(f)
    def w(*a, **k):
        if not session.get("ok"):
            if request.path.startswith("/api"):
                return jsonify(error="login"), 401
            return redirect("/")
        return f(*a, **k)
    return w


@app.get("/")
def home():
    return send_file(os.path.join(os.path.dirname(os.path.abspath(__file__)), "index.html" if session.get("ok") else "login.html"))


@app.post("/login")
def login():
    d = request.get_json(force=True)
    if d.get("u") == config.APP_USERNAME and d.get("p") == config.APP_PASSWORD:
        session["ok"] = True
        return jsonify(ok=True)
    return jsonify(ok=False), 401


@app.get("/logout")
def logout():
    session.clear()
    return redirect("/")


@app.get("/api/places")
@login_required
def api_places():
    return jsonify(places.as_json())


# ---------------------------------------------------------------- cities / towns (internet se, cache ke sath)
CACHE_FILE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "places_cache.json")
_cache = json.load(open(CACHE_FILE, encoding="utf-8")) if os.path.exists(CACHE_FILE) else {}


def _save_cache():
    with open(CACHE_FILE, "w", encoding="utf-8") as f:
        json.dump(_cache, f, ensure_ascii=False)


def _uniq(items):
    return sorted({i.strip() for i in items if i and i.strip()}, key=str.lower)


@app.get("/api/countries")
@login_required
def api_countries():
    return jsonify(places.ALL_COUNTRIES)


LAST_ERR = [""]
UA = {"User-Agent": "leadgen-tool/1.0 (contact: local use)"}


OVERPASS = ("https://overpass-api.de/api/interpreter", "https://overpass.kumi.systems/api/interpreter",
            "https://overpass.private.coffee/api/interpreter", "https://maps.mail.ru/osm/tools/overpass/api/interpreter")


def _overpass(q):
    """Free server kabhi busy (HTTP 500/429/504) hota hai: har mirror ko 2 round, beech mein thora wait."""
    for rnd in range(2):
        for url in OVERPASS:
            try:
                r = httpx.post(url, data={"data": q}, timeout=45, verify=False, headers=UA)
                if r.status_code != 200:
                    raise Exception("HTTP %d" % r.status_code)
                return [e["tags"]["name"] for e in r.json().get("elements", []) if e.get("tags", {}).get("name")]
            except Exception as e:
                LAST_ERR[0] = "%s: %s" % (type(e).__name__, str(e)[:80])
                print("overpass fail:", url, LAST_ERR[0])
        time.sleep(3 + rnd * 4)
    return None


def _geocode(city, country):
    r = httpx.get("https://nominatim.openstreetmap.org/search", params={"q": f"{city}, {country}", "format": "json", "limit": 1},
                  headers=UA, timeout=30, verify=False)
    j = r.json()
    return (float(j[0]["lat"]), float(j[0]["lon"])) if j else None


@app.get("/api/cities")
@login_required
def api_cities():
    country = request.args.get("country", "").replace('"', "")
    key = "c:" + country
    builtin = places.COUNTRIES.get(country, ([], {}))[0]
    err = ""
    if not _cache.get(key):
        names = None
        try:                                                    # 1) countriesnow
            r = httpx.post("https://countriesnow.space/api/v0.1/countries/cities", json={"country": country},
                           timeout=60, verify=False, follow_redirects=True)
            names = r.json().get("data") or None
        except Exception as e:
            err = str(e)[:120]
            print("countriesnow fail:", e)
        if not names:                                           # 2) OpenStreetMap
            names = _overpass('[out:json][timeout:80];area["name:en"="%s"]["admin_level"="2"]->.c;'
                              'node(area.c)["place"~"^(city|town)$"]["name"];out tags 5000;' % country)
        if names:
            _cache[key] = names
            _save_cache()
    got = _cache.get(key, [])
    warn = "" if got else "Internet se list nahi aayi (%s). Sirf built-in cities hain." % (err or "server jawab nahi de raha")
    return jsonify(cities=_uniq(list(builtin) + list(got)), warn=warn)


@app.get("/api/towns")
@login_required
def api_towns():
    country = request.args.get("country", "").replace('"', "")
    city = request.args.get("city", "").replace('"', "")
    key = f"t2:{country}|{city}"
    builtin = places.COUNTRIES.get(country, ([], {}))[1].get(city, [])
    if key not in _cache:
        try:
            pt = _geocode(city, country)
        except Exception as e:
            LAST_ERR[0] = "Geocode %s: %s" % (type(e).__name__, str(e)[:80])
            print(LAST_ERR[0])
            return jsonify(towns=_uniq(builtin), err=True, msg=LAST_ERR[0])
        if not pt:
            _cache[key] = []
        else:
            names = _overpass('[out:json][timeout:25];node(around:15000,%f,%f)["place"~"^(town|village|suburb)$"]["name"];out tags 250;' % pt)
            if names is None:
                return jsonify(towns=_uniq(builtin), err=True, msg=LAST_ERR[0])
            _cache[key] = [n for n in names if n.lower() != city.lower()]
        _save_cache()
    return jsonify(towns=_uniq(list(builtin) + _cache.get(key, [])))


@app.post("/api/start")
@login_required
def api_start():
    d = request.get_json(force=True)
    cats = [c for c in d.get("cats", []) if c in CATS]
    pl = [(p, p.split(",")[-1].strip()) for p in d.get("places", [])]   # "Town, City" -> (query-place, city)
    if JOB["running"] or not cats or not pl:
        return jsonify(error="Category aur kam az kam ek shehar chunein (ya search pehle se chal rahi hai)"), 400
    threading.Thread(target=run_job, args=(d.get("country", ""), pl, cats, int(d.get("max", 20))), daemon=True).start()
    return jsonify(ok=True, total=len(make_queries(d.get("country", ""), pl, cats)))


@app.post("/api/stop")
@login_required
def api_stop():
    JOB["stop"] = True
    return jsonify(ok=True)


@app.get("/api/status")
@login_required
def api_status():
    return jsonify({k: JOB[k] for k in ("running", "done", "total", "msg", "leads", "saved", "phase")})


@app.get("/download")
@login_required
def download():
    wb = Workbook()
    ws = wb.active
    ws.append(["Business", "Phone", "Email", "Website", "City", "Country", "Owner", "Title", "Facebook", "Instagram", "LinkedIn", "Twitter", "YouTube", "TikTok"])
    for l in JOB["leads"]:
        ws.append(list(l.values()))
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name="gloves_leads.xlsx")


if __name__ == "__main__":
    # Cloud (Railway/Render) ke liye PORT env variable use karo
    # Local ke liye 5055 default
    IS_CLOUD = os.environ.get("RAILWAY_ENVIRONMENT") or os.environ.get("RENDER")
    port = int(os.environ.get("PORT", 5055))

    if IS_CLOUD:
        # Cloud mode: seedha run karo, browser mat kholo
        print(f"Cloud pe chal raha hai, port: {port}")
        app.run(host="0.0.0.0", port=port, debug=False)
    else:
        # Local mode: khali port dhoondo, browser kholo
        while True:
            with socket.socket() as t:
                if t.connect_ex(("127.0.0.1", port)) != 0:
                    break
            port += 1
        url = f"http://127.0.0.1:{port}"
        print("\n>>> Browser khud khulega. Na khule to ye address copy karke Chrome mein paste karo:", url, "\n")
        threading.Timer(1.5, lambda: webbrowser.open(url)).start()
        app.run(host="127.0.0.1", port=port, debug=False)
